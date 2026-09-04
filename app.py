import streamlit as st
import pandas as pd
import openpyxl
from datetime import datetime
import io
import re

# 页面基础配置
st.set_page_config(
    page_title="OOH 投放结案 Summary 模板自动化生成器",
    page_icon="📊",
    layout="wide"
)

st.title("📊 OOH 投放结案 Summary 模板自动化生成工具")
st.write("上传 **Spotplan 表格**、**统计 DB 表格** 以及 **结案 Summary 模板文件**。系统将**自动提取主媒体折扣关联计算赠送净价价值（W列）**，完美保留动态计算公式。")

st.divider()

# 文件上传区域
col1, col2, col3 = st.columns(3)

with col1:
    spot_file = st.file_uploader("1. 上传 Spotplan (.xlsx)", type=["xlsx"], key="spot")

with col2:
    db_file = st.file_uploader("2. 上传 统计 DB (.xlsx)", type=["xlsx"], key="db")

with col3:
    template_file = st.file_uploader("3. 上传 Summary 模板 (.xlsx)", type=["xlsx"], key="template")

def clean_location_name(loc_name):
    """提取并清洗媒体名称，去掉（赠送/额外赠送/增值）字眼"""
    if not loc_name:
        return "", "", ""
    loc = str(loc_name).strip()
    loc_clean = re.sub(r'[（\(](赠送|额外赠送|增值)[）\)]', '', loc).strip()
    loc_clean2 = re.sub(r'[（\(]\d+块/套[）\)]', '', loc_clean).strip()
    return loc_clean, loc_clean2, loc

def parse_days_from_period(period_str):
    """解析如 2026.03.01-2026.03.14 的周期格式并计算天数"""
    if not period_str:
        return 0
    match = re.search(r'(\d{4}[\.\/-]\d{1,2}[\.\/-]\d{1,2})\s*[-~至]\s*(\d{4}[\.\/-]\d{1,2}[\.\/-]\d{1,2})', str(period_str))
    if match:
        try:
            d1_str = re.sub(r'[\/-]', '.', match.group(1))
            d2_str = re.sub(r'[\/-]', '.', match.group(2))
            d1 = datetime.strptime(d1_str, "%Y.%m.%d")
            d2 = datetime.strptime(d2_str, "%Y.%m.%d")
            return (d2 - d1).days + 1
        except Exception:
            return 0
    return 0

def generate_summary_from_template(spot_file, db_file, template_file):
    # -------------------------------------------------------------
    # 1. 读取 统计 DB 构建 Lookup 检索字典
    # -------------------------------------------------------------
    wb_db = openpyxl.load_workbook(db_file, data_only=True)
    ws_db = wb_db.active
    if '统计DB' in wb_db.sheetnames:
        ws_db = wb_db['统计DB']
    
    db_lookup = {}
    for r in range(2, ws_db.max_row + 1):
        loc_val = ws_db.cell(r, 13).value       # Col M: Location
        coverage_val = ws_db.cell(r, 40).value  # Col AN: 有效覆盖人车次/天（单块）
        
        if loc_val is not None:
            loc_str = str(loc_val).strip()
            if loc_str and loc_str not in db_lookup:
                db_lookup[loc_str] = coverage_val
            c1, c2, _ = clean_location_name(loc_str)
            if c1 and c1 not in db_lookup:
                db_lookup[c1] = coverage_val
            if c2 and c2 not in db_lookup:
                db_lookup[c2] = coverage_val

    # -------------------------------------------------------------
    # 2. 读取 Spotplan 明细数据
    # -------------------------------------------------------------
    wb_spot = openpyxl.load_workbook(spot_file, data_only=True)
    ws_spot = wb_spot.active
    
    header_row = 4
    for r in range(1, 10):
        row_vals = [str(ws_spot.cell(r, c).value or '') for c in range(1, 15)]
        if 'Market' in row_vals and 'Location' in row_vals:
            header_row = r
            break
            
    spot_rows = []
    for r in range(header_row + 1, ws_spot.max_row + 1):
        mkt = ws_spot.cell(r, 2).value  # Col B
        loc = ws_spot.cell(r, 4).value  # Col D
        if mkt or loc:
            row_data = [ws_spot.cell(r, c).value for c in range(2, 18)]
            spot_rows.append(row_data)

    # -------------------------------------------------------------
    # 3. 加载模板并预处理：扫描无“赠送”字样的主媒体折扣与对应行号
    # -------------------------------------------------------------
    wb_tpl = openpyxl.load_workbook(template_file)
    ws_tpl = wb_tpl.active
    
    start_row = 4  # 数据写入起始行
    
    # 取消数据填入区域的合并单元格
    merged_ranges = list(ws_tpl.merged_cells.ranges)
    for rng in merged_ranges:
        if rng.max_row >= start_row:
            ws_tpl.unmerge_cells(str(rng))

    # 【核心新增步骤】：映射主媒体与其对应的“行号”及“折扣值”
    main_media_info = {} # key: cleaned_name, value: {'row': current_row, 'discount': discount_val}
    for idx, row in enumerate(spot_rows):
        current_row = start_row + idx
        loc = str(row[2] or '')
        c1, c2, orig = clean_location_name(loc)
        is_bonus = ("赠送" in orig) or ("额外赠送" in orig) or ("增值" in orig)
        
        # 如果是不带赠送的主媒体，记录其在输出表中的Excel行号和折扣
        if not is_bonus and c1:
            discount_val = row[10] # Col L: Discount
            main_media_info[c1] = {'row': current_row, 'discount': discount_val}
            if c2:
                main_media_info[c2] = {'row': current_row, 'discount': discount_val}

    # 提取模板中第 4 行的样式作为样式基准
    sample_cells = [ws_tpl.cell(start_row, col) for col in range(1, 30)]

    # -------------------------------------------------------------
    # 4. 填充数据、匹配主媒体折扣并嵌入原生 Excel 计算公式
    # -------------------------------------------------------------
    for idx, row in enumerate(spot_rows):
        current_row = start_row + idx
        
        mkt = row[0]          # Col B: Market
        fmt = row[1]          # Col C: Format
        loc = str(row[2] or '')# Col D: Location
        period = row[3]       # Col E: 计划投放周期
        no_week = row[4]      # Col F: No. Of Week
        no_unit = row[5]      # Col G: No. Of Unit
        buying_unit = row[6]  # Col H: Buying Uint
        duration_freq = row[7]# Col I: Duration/Frequency
        ratecard_cost = row[8]# Col J: Ratecard Cost
        ratecard_ttl = row[9] # Col K: Ratecard TTL Cost
        discount = row[10]    # Col L: Discount
        net_unit_cost = row[11]# Col M: Net Unit Cost
        net_ttl_cost = row[12] # Col N: Net TTL Cost
        unit_prod_fee = row[13]# Col O: Unit Production Fee
        prod_fee = row[14]    # Col P: Production Fee
        gross_cost = row[15]  # Col Q: Gross Cost

        # 判定是否为“赠送”点位
        c1, c2, orig = clean_location_name(loc)
        is_bonus = ("赠送" in orig) or ("额外赠送" in orig) or ("增值" in orig)
        
        bonus_period_text = ""
        actual_period_text = period
        
        if is_bonus:
            bonus_period_text = f"赠送周期: {period}"
            actual_period_text = f"{period} (赠送)"

        # ---------------------------------------------------------
        # 【核心计算】：为“赠送”点位寻找主媒体的折扣与关联公式
        # ---------------------------------------------------------
        w_net_value_formula = 0
        if is_bonus:
            # 尝试在主媒体映射表中寻找对应的无赠送字样点位
            matched_main = main_media_info.get(c1) or main_media_info.get(c2)
            if matched_main:
                main_row = matched_main['row']
                # 植入原生 Excel 公式：赠送行K列 (刊例) * 主媒体行L列 (折扣)
                w_net_value_formula = f"=K{current_row}*L{main_row}"
            else:
                # 若未查到对应的主媒体行，则尝试直接乘以本行的折扣
                w_net_value_formula = f"=K{current_row}*L{current_row}"

        # C列 资源数量
        resource_qty = f"{no_unit}{buying_unit}" if (no_unit and buying_unit) else no_unit

        # AA列 日均覆盖人次检索
        daily_coverage = db_lookup.get(orig) or db_lookup.get(c1) or db_lookup.get(c2)
        if daily_coverage == "/" or daily_coverage is None:
            for k, v in db_lookup.items():
                if c2 and c2 in k and v != "/":
                    daily_coverage = v
                    break

        # 构造包含 Excel 原生计算公式的列映射结构
        row_values = {
            1: mkt,                                           # A: 市场
            2: loc,                                           # B: 媒体
            3: resource_qty,                                  # C: 资源数量
            4: duration_freq,                                 # D: 广告频次
            5: period,                                        # E: 计划投放周期
            6: no_week,                                       # F: No. Of Week
            7: no_unit,                                       # G: No. Of Unit
            8: buying_unit,                                   # H: Buying Uint
            9: duration_freq,                                 # I: Duration/Frequency
            10: ratecard_cost,                                # J: Ratecard Cost
            11: f"=F{current_row}*G{current_row}*J{current_row}", # K: Ratecard TTL Cost 公式
            12: discount,                                     # L: Discount
            13: f"=J{current_row}*L{current_row}",             # M: Net Unit Cost 公式
            14: f"=F{current_row}*G{current_row}*M{current_row}", # N: Net TTL Cost 公式
            15: unit_prod_fee,                                # O: Unit Production Fee
            16: f"=G{current_row}*O{current_row}",             # P: Production Fee 公式
            17: f"=N{current_row}+P{current_row}",             # Q: Gross Cost 公式
            18: bonus_period_text if is_bonus else "",        # R: 额外赠送（显示赠送周期）
            19: actual_period_text,                           # S: 实际投放周期
            20: f"=N{current_row}",                            # T: 投放实际总净价 公式
            21: f"=P{current_row}",                            # U: 投放实际总制作费 公式
            22: f"=K{current_row}" if is_bonus else 0,         # V: 投放赠送价值（刊例直接等于自身K列）
            23: w_net_value_formula if is_bonus else 0,       # W: 投放赠送价值（净价 = K列 * 主媒体折扣L列）
            24: 0,                                            # X: 补偿价值（净价）
            25: 0,                                            # Y: 非补偿增值赠送（刊例）
            26: 0,                                            # Z: 非补偿增值赠送（净价）
            27: daily_coverage if daily_coverage is not None else 0, # AA: 日均覆盖人次
            28: f"=F{current_row}*7",                         # AB: 投放天数 公式 (周数*7)
            29: f"=AA{current_row}*AB{current_row}"           # AC: 总覆盖人次 公式 (日均*天数)
        }

        # 写入单元格并完全复刻模板样式
        for col_idx, val in row_values.items():
            cell = ws_tpl.cell(row=current_row, column=col_idx)
            cell.value = val
            
            # 继承模板第 4 行单元格样式
            sample_cell = sample_cells[col_idx - 1]
            if sample_cell.has_style:
                cell.font = sample_cell.font.copy()
                cell.fill = sample_cell.fill.copy()
                cell.border = sample_cell.border.copy()
                cell.alignment = sample_cell.alignment.copy()
                cell.number_format = sample_cell.number_format

    # 导出文件字节流
    output = io.BytesIO()
    wb_tpl.save(output)
    output.seek(0)
    return output

# 触发按钮逻辑
if spot_file and db_file and template_file:
    if st.button("🚀 套用模板生成结案 Summary (关联主媒体折扣与公式)", type="primary"):
        try:
            with st.spinner("正在联动主媒体折扣，计算赠送价值并写入 Excel 公式..."):
                excel_out = generate_summary_from_template(spot_file, db_file, template_file)
                st.success("🎉 生成成功！已成功匹配主点位折扣，V列与W列已嵌入跨行动态关联公式。")
                st.download_button(
                    label="📥 点击下载模板渲染 Summary.xlsx",
                    data=excel_out,
                    file_name="OOH_投放结案_Summary_主折扣联动版.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
        except Exception as e:
            st.error(f"生成失败，错误原因: {str(e)}")
else:
    st.info("💡 请在上方的 3 个上传框中分别拖入：Spotplan 表格、统计 DB 表格 和 Summary 模板表格。")